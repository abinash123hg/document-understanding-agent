import os
from pathlib import Path
from typing import Optional
import requests
import fitz
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

try:
    from docx import Document
except Exception:
    Document = None

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

MODEL = os.getenv("LLM_MODEL", "qwen2.5:1.5b")
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")

app = FastAPI(title="Private Document Chat")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:5500", "http://localhost:5500", "*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ChatRequest(BaseModel):
    question: str
    documentname: Optional[str] = None

def safe_document_path(name: Optional[str]) -> Path:
    if name:
        clean_name = Path(name).name
        if clean_name != name or clean_name in ("", ".", ".."):
            raise HTTPException(400, "Invalid document name.")
        path = UPLOAD_DIR / clean_name
        if not path.is_file():
            raise HTTPException(404, "Selected document was not found.")
        return path

    files = sorted(
        [p for p in UPLOAD_DIR.iterdir() if p.is_file()],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        raise HTTPException(404, "No document uploaded.")
    return files[0]

def read_docx(path: Path) -> str:
    if Document is None:
        raise HTTPException(500, "python-docx is not installed.")

    doc = Document(str(path))
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)

def read_pdf(path: Path) -> str:
    pages = []

    with fitz.open(str(path)) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(f"[Page {page_number}]\n{text}")

    return "\n\n".join(pages)

def read_document(path: Path) -> str:
    extension = path.suffix.lower()

    if extension == ".docx":
        return read_docx(path)

    if extension == ".pdf":
        return read_pdf(path)

    if extension in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    raise HTTPException(
        400,
        "Unsupported file type. Use PDF, DOCX, TXT, or MD."
    )

def build_prompt(document_text: str, question: str) -> str:
    limited_text = document_text[:120000]

    return f"""
You are a document-only Q&A system.

STRICT RULES:
1. Answer ONLY from the DOCUMENT below.
2. Do NOT use outside knowledge.
3. Do NOT invent facts, numbers, tables, or examples.
4. If the answer is not in the document, reply EXACTLY:
I could not find the answer in this document.
5. Do NOT create tables unless the document contains tables.
6. Do NOT guess customer IDs, names, ages, or amounts.

DOCUMENT:
--- BEGIN DOCUMENT ---
{limited_text}
--- END DOCUMENT ---

USER QUESTION:
{question}
""".strip()

def ask_ollama(prompt: str) -> str:
    payload = {
        "model": MODEL,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You answer only from supplied document text. "
                    "When evidence is missing, use the exact fallback sentence."
                ),
            },
            {
                "role": "user",
                "content": prompt,
            },
        ],
        "stream": False,
        "options": {
            "temperature": 0,
            "top_p": 0.1,
            "repeat_penalty": 1.15,
            "num_ctx": 32768,
            "num_predict": 700,
        },
    }

    try:
        response = requests.post(
            f"{OLLAMA_URL}/api/chat",
            json=payload,
            timeout=180,
        )
        response.raise_for_status()
        data = response.json()
    except requests.exceptions.ConnectionError:
        raise HTTPException(
            503,
            "Ollama is not running. Start Ollama and try again."
        )
    except requests.exceptions.Timeout:
        raise HTTPException(504, "Ollama response timed out.")
    except requests.RequestException as error:
        raise HTTPException(502, f"Ollama request failed: {error}")
    except ValueError:
        raise HTTPException(502, "Ollama returned invalid JSON.")

    answer = (
        data.get("message", {}).get("content", "")
        if isinstance(data, dict)
        else ""
    )

    answer = str(answer).strip()

    if not answer:
        raise HTTPException(502, "Ollama returned an empty answer.")

    return answer

@app.get("/")
def root():
    return {
        "status": "running",
        "model": MODEL,
        "document_mode": True,
    }

@app.get("/health")
def health():
    try:
        response = requests.get(f"{OLLAMA_URL}/api/tags", timeout=5)
        return {
            "backend": "ok",
            "ollama": response.ok,
            "model": MODEL,
        }
    except Exception as error:
        return {
            "backend": "ok",
            "ollama": False,
            "model": MODEL,
            "error": str(error),
        }

@app.post("/upload")
async def upload(file: UploadFile = File(...)):
    original_name = Path(file.filename or "document.txt").name
    extension = Path(original_name).suffix.lower()

    if extension not in {".pdf", ".docx", ".txt", ".md"}:
        raise HTTPException(
            400,
            "Unsupported file type. Use PDF, DOCX, TXT, or MD."
        )

    content = await file.read()

    if not content:
        raise HTTPException(400, "Uploaded file is empty.")

    destination = UPLOAD_DIR / original_name
    destination.write_bytes(content)

    text = read_document(destination)

    if not text.strip():
        raise HTTPException(
            400,
            "No readable text was found in this document."
        )

    return {
        "filename": original_name,
        "documentname": original_name,
        "characters": len(text),
    }

@app.post("/documents/upload")
async def documents_upload(file: UploadFile = File(...)):
    return await upload(file)

@app.post("/chat")
def chat(request: ChatRequest):
    question = request.question.strip()

    if not question:
        raise HTTPException(400, "Question cannot be empty.")

    document_path = safe_document_path(request.documentname)
    document_text = read_document(document_path).strip()

    if not document_text:
        raise HTTPException(
            400,
            "No readable text was found in the selected document."
        )

    prompt = build_prompt(document_text, question)
    answer = ask_ollama(prompt)

    return {
        "answer": answer,
        "response": answer,
        "documentname": document_path.name,
        "sources": [document_path.name],
    }


