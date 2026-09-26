import os
from pathlib import Path
from typing import Optional

import requests
import fitz
import pytesseract

from PIL import Image
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

try:
    from docx import Document
except Exception:
    Document = None

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

app = FastAPI(title="Fast Private Document Chat")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:5500",
        "http://localhost:5500",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MODEL = os.getenv("LLM_MODEL", "qwen2.5:1.5b")
OLLAMA_URL = "http://127.0.0.1:11434/api/chat"
UPLOAD_DIR = Path(__file__).resolve().parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

class ChatRequest(BaseModel):
    question: str
    document_name: Optional[str] = None

def read_docx(path: Path) -> str:
    if Document is None:
        raise HTTPException(500, "python-docx package missing")

    doc = Document(str(path))
    text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n".join(text)

def read_pdf(path: Path) -> str:
    document = fitz.open(str(path))
    pages = []

    for page in document:
        page_text = page.get_text("text").strip()

        if page_text:
            pages.append(page_text)
            continue

        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(1.5, 1.5),
            alpha=False
        )

        image = Image.frombytes(
            "RGB",
            [pixmap.width, pixmap.height],
            pixmap.samples
        )

        ocr_text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        ).strip()

        if ocr_text:
            pages.append(ocr_text)

    return "\n".join(pages)

def read_document(path: Path) -> str:
    extension = path.suffix.lower()

    if extension == ".docx":
        return read_docx(path)

    if extension == ".pdf":
        return read_pdf(path)

    if extension in [".txt", ".md"]:
        return path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

    raise HTTPException(
        400,
        f"Unsupported file type: {extension}"
    )

def get_document(name: Optional[str]) -> Path:
    if name:
        requested = UPLOAD_DIR / Path(name).name
        if requested.exists():
            return requested

    files = sorted(
        [file for file in UPLOAD_DIR.iterdir() if file.is_file()],
        key=lambda file: file.stat().st_mtime,
        reverse=True
    )

    if not files:
        raise HTTPException(
            404,
            "No document uploaded"
        )

    return files[0]

@app.get("/")
def root():
    return {
        "status": "running",
        "model": MODEL,
        "ocr": True
    }

@app.get("/health")
def health():
    try:
        response = requests.get(
            "http://127.0.0.1:11434/api/tags",
            timeout=5
        )

        return {
            "backend": "ok",
            "ollama": response.ok,
            "model": MODEL,
            "ocr": True
        }

    except Exception as error:
        return {
            "backend": "ok",
            "ollama": False,
            "model": MODEL,
            "ocr": True,
            "error": str(error)
        }

@app.post("/upload")
async def upload(file: UploadFile = File(...)):
    filename = Path(
        file.filename or "document"
    ).name

    destination = UPLOAD_DIR / filename
    destination.write_bytes(await file.read())

    return {
        "filename": filename,
        "document_name": filename
    }

@app.post("/documents/upload")
async def documents_upload(file: UploadFile = File(...)):
    return await upload(file)

@app.post("/chat")
def chat(request: ChatRequest):
    document_path = get_document(request.document_name)
    document_text = read_document(document_path)

    if not document_text.strip():
        raise HTTPException(
            400,
            "Document se text nahi mila."
        )

    prompt = f"""You are a strict private document assistant.

LANGUAGE:
LANGUAGE DETECTION:
- Detect the language of the USER QUESTION yourself.
- If the user writes English, reply completely in English.
- If the user writes Hindi in Devanagari, reply completely in Hindi.
- If the user writes Hinglish using English letters, reply in Hinglish.
- Do not change language unless the user asks you to.
- Never reply with Hindi/Hinglish when the user asked in English.

DOCUMENT RULES:
- For document questions, use only the document below.
- Never invent facts, file names, code, features, people, or sections.
- Never guess.
- If the answer is not clearly present, reply in the user's detected language:
- English: "I could not find the answer in this document."
- Hindi: "मुझे इस दस्तावेज़ में इसका उत्तर नहीं मिला।"
- Hinglish: "Is document mein iska answer nahi mila."
- Answer completely but clearly. Do not stop mid-sentence. For detailed questions, use bullets and finish every bullet.
- For summaries, include only information found in the document.
- For code questions, mention only code actually present in the document.
- General questions may be answered as general information.

DOCUMENT:
{document_text[:18000]}

USER QUESTION:
{request.question}
"""

    payload = {
        "model": MODEL,
        "messages": [
            {
                "role": "user",
                "content": prompt
            }
        ],
        "stream": False,
        "keep_alive": "30m",
        "options": {
            "temperature": 0.0,
            "num_ctx": 2048,
            "num_predict": 700,
            "top_p": 0.2,
            "repeat_penalty": 1.1
        }
    }

    try:
        response = requests.post(
            OLLAMA_URL,
            json=payload,
            timeout=180
        )
    except Exception as error:
        raise HTTPException(
            503,
            f"Ollama connect nahi ho raha: {error}"
        )

    if not response.ok:
        raise HTTPException(
            response.status_code,
            f"Ollama error: {response.text}"
        )

    try:
        result = response.json()
        answer = result.get(
            "message",
            {}
        ).get(
            "content",
            ""
        ).strip()
    except Exception as error:
        raise HTTPException(
            502,
            f"Invalid Ollama response: {error}"
        )

    if not answer:
        raise HTTPException(
            502,
            "Ollama ne empty answer diya"
        )

    return {
        "answer": answer,
        "response": answer,
        "document_name": document_path.name
    }

